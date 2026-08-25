"""Utilidades de formato para los documentos Word del proyecto.

Compartidas por hacer_resumen.py y hacer_muestra.py. Se extrajeron cuando el
segundo documento iba a copiar las mismas sesenta lineas: dos copias del formato
se desincronizan a la primera correccion.
"""

import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent / ".libs"))

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

TINTA = RGBColor(0x1A, 0x23, 0x31)      # azul-negro de imprenta
ROJO = RGBColor(0xB2, 0x2F, 0x26)       # rojo de corrector: marca lo que falla
GRIS = RGBColor(0x5A, 0x63, 0x6E)
VERDE = RGBColor(0x2C, 0x6A, 0x4E)


def sombrear(celda, hex_color):
    tinte = OxmlElement("w:shd")
    tinte.set(qn("w:val"), "clear")
    tinte.set(qn("w:fill"), hex_color)
    celda._tc.get_or_add_tcPr().append(tinte)


def linea_inferior(parrafo, color="1A2331", grosor=8):
    pPr = parrafo._p.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    borde = OxmlElement("w:bottom")
    borde.set(qn("w:val"), "single")
    borde.set(qn("w:sz"), str(grosor))
    borde.set(qn("w:space"), "4")
    borde.set(qn("w:color"), color)
    bordes.append(borde)
    pPr.append(bordes)
    return parrafo


def parrafo(doc, texto="", tam=10.5, negrita=False, color=TINTA, fuente="Calibri",
            espacio_antes=0, espacio_despues=6, cursiva=False, sangria=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(espacio_antes)
    p.paragraph_format.space_after = Pt(espacio_despues)
    p.paragraph_format.line_spacing = 1.15
    if sangria:
        p.paragraph_format.left_indent = Cm(sangria)
    if texto:
        r = p.add_run(texto)
        r.font.size = Pt(tam)
        r.font.bold = negrita
        r.font.italic = cursiva
        r.font.color.rgb = color
        r.font.name = fuente
    return p


def titulo_seccion(doc, texto):
    p = parrafo(doc, texto, tam=13, negrita=True, espacio_antes=16, espacio_despues=8)
    linea_inferior(p)
    return p


def viñeta(doc, texto, negritas=()):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    _escribir_con_negritas(p, texto, negritas)
    return p


def _escribir_con_negritas(p, texto, negritas):
    """Escribe el texto resaltando en negrita los fragmentos indicados."""
    resto = texto
    while resto:
        pos, encontrado = len(resto), None
        for frag in negritas:
            i = resto.find(frag)
            if i != -1 and i < pos:
                pos, encontrado = i, frag
        if encontrado is None:
            r = p.add_run(resto)
            r.font.size = Pt(10.5)
            r.font.color.rgb = TINTA
            break
        if pos:
            r = p.add_run(resto[:pos])
            r.font.size = Pt(10.5)
            r.font.color.rgb = TINTA
        r = p.add_run(encontrado)
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = TINTA
        resto = resto[pos + len(encontrado):]


def tabla(doc, cabeceras, filas, anchos):
    t = doc.add_table(rows=1, cols=len(cabeceras))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (celda, ancho) in enumerate(zip(t.rows[0].cells, anchos)):
        celda.width = Cm(ancho)
        sombrear(celda, "1A2331")
        celda.paragraphs[0].paragraph_format.space_after = Pt(2)
        r = celda.paragraphs[0].add_run(cabeceras[i])
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].width = Cm(anchos[i])
            celdas[i].paragraphs[0].paragraph_format.space_after = Pt(2)
            r = celdas[i].paragraphs[0].add_run(str(valor))
            r.font.size = Pt(9.5)
            r.font.color.rgb = TINTA
    return t


def evidencia(doc, lineas):
    """Bloque de salida real del auditor, en monoespaciada."""
    caja = doc.add_table(rows=1, cols=1)
    caja.style = "Table Grid"
    celda = caja.rows[0].cells[0]
    celda.width = Cm(16.5)
    sombrear(celda, "F4F3F0")
    celda.paragraphs[0]._p.getparent().remove(celda.paragraphs[0]._p)
    for texto, color in lineas:
        p = celda.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        r = p.add_run(texto)
        r.font.size = Pt(8.5)
        r.font.name = "Consolas"
        r.font.color.rgb = color
    return caja
