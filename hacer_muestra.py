r"""Arma el Word con las piezas de muestra para que las evalue Edicion.

    .pyruntime\python.exe hacer_muestra.py

Lee lo que hay en borradores/ y lo maqueta. No reescribe nada: lo que Pablo lee
es exactamente lo que produjo el sistema, con sus aciertos y sus fallos.
"""

import json
import pathlib
import sys

AQUI = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI / ".libs"))
sys.path.insert(0, str(AQUI))

from docx import Document
from docx.shared import Cm, Pt

from docx_util import (GRIS, ROJO, TINTA, VERDE, linea_inferior, parrafo,
                       sombrear, tabla, titulo_seccion, viñeta)

PIEZAS = [
    ("gtm_remesas_noticia", "Pieza 1 · Noticia",
     "Pone a prueba la regla difícil: el cuerpo va sin postura y la postura vive "
     "solo en el bloque «SurEconomics:»."),
    ("ven_inflacion_educación", "Pieza 2 · Artículo educativo",
     "Pone a prueba el registro para persona normal y el manejo de un dato viejo "
     "(la última inflación que publica el Banco Mundial para Venezuela es de 2016)."),
]


def caja_pieza(doc, texto):
    """La pieza tal cual, en un recuadro, para que se lea como pieza y no como informe."""
    caja = doc.add_table(rows=1, cols=1)
    caja.style = "Table Grid"
    celda = caja.rows[0].cells[0]
    celda.width = Cm(16.5)
    sombrear(celda, "FBFAF8")
    celda.paragraphs[0]._p.getparent().remove(celda.paragraphs[0]._p)

    for i, linea in enumerate(texto.split("\n")):
        p = celda.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if not linea.strip():
            continue
        r = p.add_run(linea)
        if linea.startswith("["):          # etiquetas y metadatos
            r.font.size, r.font.color.rgb, r.font.bold = Pt(8.5), GRIS, True
        elif linea.isupper() and len(linea) > 25:   # titular
            r.font.size, r.font.bold, r.font.color.rgb = Pt(13), True, TINTA
        elif linea.startswith("SurEconomics:"):
            r.font.size, r.font.bold, r.font.color.rgb = Pt(10.5), True, VERDE
        elif linea.startswith("—"):
            r.font.size, r.font.color.rgb = Pt(8.5), GRIS
        else:
            r.font.size, r.font.color.rgb = Pt(10.5), TINTA
    return caja


doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.4)
estilo = doc.styles["Normal"]
estilo.font.name = "Calibri"
estilo.font.size = Pt(10.5)
estilo.font.color.rgb = TINTA

parrafo(doc, "SURECONOMICS · PRUEBA DE REDACCIÓN AUTOMATIZADA", tam=8.5,
        negrita=True, color=GRIS, espacio_despues=2)
p = parrafo(doc, "Dos piezas para evaluación editorial", tam=19, negrita=True,
            espacio_despues=4)
p.paragraph_format.line_spacing = 1.0
linea_inferior(parrafo(doc, "Para Pablo Quintero · Jefatura Editorial · "
                            "24 de agosto de 2026", tam=9.5, color=GRIS,
                       espacio_despues=10))

titulo_seccion(doc, "Qué estás leyendo")
parrafo(doc,
        "Estas dos piezas las escribió el sistema de punta a punta, sin que nadie "
        "las tocara después. El proceso fue: el código fue al Banco Mundial y trajo "
        "las cifras con su enlace; la IA escribió la prosa usando solo esas cifras; "
        "y un auditor verificó el texto contra ellas antes de dejarlo pasar.",
        espacio_despues=6)
parrafo(doc,
        "Las dos pasaron el auditor. Eso significa que no hay ninguna cifra "
        "inventada, ninguna fuente falsa y ningún error de formato — pero NO "
        "significa que estén bien escritas. Eso es justamente lo que te pedimos "
        "que juzgues.", espacio_despues=6)
parrafo(doc, "El sistema no publica nada. Todo entra como borrador.",
        negrita=True, color=VERDE)

titulo_seccion(doc, "Qué nos sirve que evalúes")
viñeta(doc, "El registro. ¿Está escrito para una persona normal o todavía suena a "
            "informe? ¿Responde a «esto a mí en qué me afecta»?",
       negritas=["El registro."])
viñeta(doc, "La separación hecho / postura en la noticia. ¿El cuerpo quedó "
            "realmente limpio? ¿El bloque «SurEconomics:» dice lo que debe decir y "
            "con la fuerza que debe tener?",
       negritas=["La separación hecho / postura en la noticia."])
viñeta(doc, "Los titulares. ¿Sirven? ¿Prometen más de lo que el cuerpo sostiene?",
       negritas=["Los titulares."])
viñeta(doc, "La línea editorial. ¿El bloque de la noticia refleja la posición del "
            "medio, o se queda tibio?",
       negritas=["La línea editorial."])
viñeta(doc, "Las etiquetas. ¿Clasificó bien tipo, región, subregión, país y tópico?",
       negritas=["Las etiquetas."])

titulo_seccion(doc, "Lo que ya vimos nosotros y no arreglamos a propósito")
parrafo(doc, "Para que la muestra sea honesta, dejamos los fallos que detectamos:",
        espacio_despues=6)
tabla(doc, ["Dónde", "Qué pasa"],
      [["Pieza 2, último párrafo",
        "Escribe «Para el lector normal…». Es lenguaje de la instrucción que se "
        "coló en el texto. Hay que corregirlo en el prompt."],
       ["Pieza 1, segundo párrafo",
        "Dice que las remesas «reflejan la dependencia de la economía». Puede ser "
        "postura dentro del cuerpo, que es justo lo que no debe pasar. Tu criterio "
        "decide: si lo es, ajustamos el prompt."],
       ["Pieza 1, bloque final",
        "Dice «SurEconomics aboga por…» en tercera persona. El perfil pide primera "
        "persona plural."]],
      [4.5, 12.0])

for archivo, titulo, proposito in PIEZAS:
    doc.add_page_break()
    titulo_seccion(doc, titulo)
    parrafo(doc, proposito, tam=9.5, color=GRIS, cursiva=True, espacio_despues=8)
    ruta = AQUI / "borradores" / f"{archivo}.txt"
    caja_pieza(doc, ruta.read_text(encoding="utf-8").strip())

    datos = json.loads((AQUI / "borradores" / f"{archivo}.json").read_text(encoding="utf-8"))
    parrafo(doc, "Veredicto del auditor", negrita=True, tam=10, espacio_antes=10,
            espacio_despues=4)
    estado = "BLOQUEADA" if datos["bloqueada"] else "APROBADA para la cola editorial"
    parrafo(doc, estado, negrita=True, color=ROJO if datos["bloqueada"] else VERDE,
            tam=10, espacio_despues=4)
    for h in datos["hallazgos"]:
        parrafo(doc, h.strip(), tam=9, color=GRIS, sangria=0.5, espacio_despues=2)

salida = AQUI / "Piezas de muestra - SurEconomics.docx"
doc.save(salida)
print("Guardado:", salida)
